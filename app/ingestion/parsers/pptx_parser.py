"""pptx 파서: 슬라이드를 단위로, 텍스트 프레임과 표를 추출.

슬라이드 내 이미지/도표는 필요 시 문서 파서(PaddleOCR-VL)로 캡션을 생성하도록
image_ocr 콜백을 주입할 수 있다(없으면 이미지 캡션은 생략).
"""

from __future__ import annotations

from typing import Callable, Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.schemas.enums import ChunkType, FileFormat

from .base import ParsedElement, ParseResult

ImageOCR = Callable[[bytes], str]


def _table_to_markdown(table) -> str:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    header, *body = rows
    lines = ["| " + " | ".join(header) + " |",
             "| " + " | ".join("---" for _ in header) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


class PptxParser:
    file_format = FileFormat.PPTX

    def __init__(self, image_ocr: Optional[ImageOCR] = None) -> None:
        self._image_ocr = image_ocr

    def parse(self, path: str) -> ParseResult:
        prs = Presentation(path)
        elements: list[ParsedElement] = []
        slides = list(prs.slides)

        for idx, slide in enumerate(slides, start=1):
            section = f"슬라이드 {idx}"
            for shape in slide.shapes:
                if shape.has_table:
                    md = _table_to_markdown(shape.table)
                    if md:
                        elements.append(ParsedElement(
                            text=md, element_type=ChunkType.TABLE,
                            section_title=section, page_no=idx))
                elif shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        elements.append(ParsedElement(
                            text=text, element_type=ChunkType.TEXT,
                            section_title=section, page_no=idx))
                elif (self._image_ocr is not None
                      and shape.shape_type == MSO_SHAPE_TYPE.PICTURE):
                    try:
                        caption = self._image_ocr(shape.image.blob).strip()
                    except Exception:
                        caption = ""
                    if caption:
                        elements.append(ParsedElement(
                            text=caption, element_type=ChunkType.IMAGE_CAPTION,
                            section_title=section, page_no=idx))

        return ParseResult(elements=elements, page_count=len(slides))
