"""표 데이터 적재(Tier 2): 확정된 xlsx 원본 → DuckDB 테이블 + 카탈로그.

임계값을 넘는 시트만 구조화 적재한다(작은 표는 Tier 1 임베딩으로 충분).
메모리 안전: 미리보기(1차)로 스키마·타입을 정하고, 2차에서 스트리밍 배치 삽입.
"""

from __future__ import annotations

import os

from sqlalchemy.orm import Session

from app.datasets.detect import (
    DATA_TABLE_ROW_THRESHOLD,
    SAMPLE_ROWS,
    is_data_table,
    normalize_headers,
    read_sheet_preview,
)
from app.datasets.store import DuckDBStore, infer_types, safe_table_name
from app.db.repositories import DatasetRepository, DocumentRepository
from app.schemas.enums import FileFormat


def _data_rows(ws):
    """헤더 제외, 빈 행 제외한 데이터행을 스트리밍(메모리 절약)."""
    first = True
    for row in ws.iter_rows(values_only=True):
        if row is None:
            continue
        cells = ["" if c is None else str(c).strip() for c in row]
        if not any(cells):
            continue
        if first:
            first = False   # 헤더 스킵
            continue
        yield cells


def ingest_if_tabular(session: Session, doc, store: DuckDBStore | None = None) -> int:
    """문서가 xlsx 이고 대용량 표 시트를 포함하면 DuckDB 에 적재. 반환: 만든 데이터셋 수."""
    if doc.identification.file_format != FileFormat.XLSX:
        return 0
    doc_id = doc.identification.doc_id
    path = DocumentRepository(session).get_original_path(doc_id)
    if not path or not os.path.exists(path):
        return 0

    import openpyxl
    store = store or DuckDBStore()
    catalog = DatasetRepository(session)
    tokens = list(doc.governance.access_tokens) or ["*"]

    # 재적재: 기존 데이터셋/테이블 정리
    for old in catalog.drop_for_doc(doc_id):
        store.drop(old)

    # 1차: 시트별 스키마·행수·샘플
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    plans = []
    for ws in wb.worksheets:
        header, _small, n_rows, samples = read_sheet_preview(ws, DATA_TABLE_ROW_THRESHOLD)
        if header is None or not is_data_table(n_rows):
            continue
        cols = normalize_headers(header)
        types = infer_types(cols, samples[:SAMPLE_ROWS])
        plans.append((ws.title, cols, types))
    wb.close()
    if not plans:
        return 0

    # 2차: 시트별 전체 행 스트리밍 삽입
    wb2 = openpyxl.load_workbook(path, read_only=True, data_only=True)
    created = 0
    try:
        for sheet, cols, types in plans:
            ws = wb2[sheet]
            table = safe_table_name(doc_id, sheet)
            n = store.create_from_rows(table, cols, types, _data_rows(ws))
            catalog.upsert(doc_id, sheet, table,
                           [{"name": c, "type": t} for c, t in zip(cols, types)],
                           n, tokens)
            created += 1
    finally:
        wb2.close()
    return created


def drop_datasets(session: Session, doc_id: str, store: DuckDBStore | None = None) -> None:
    """문서 삭제/대체 시 데이터셋(카탈로그 + DuckDB 테이블) 정리."""
    store = store or DuckDBStore()
    for table in DatasetRepository(session).drop_for_doc(doc_id):
        store.drop(table)
