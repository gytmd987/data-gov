"""docx 파서: 본문 텍스트·표 + 삽입 이미지 OCR.

워드에 캡처·스캔본을 붙여 넣은 문서가 흔한데, 이미지를 건너뛰면 그 내용이
검색에 전혀 안 잡힌다(같은 내용을 PDF로 올리면 잡히는데도).
"""

import docx
import pytest

from app.ingestion.parsers import get_parser
from app.ingestion.parsers.docx_parser import DocxParser
from app.schemas.enums import ChunkType, FileFormat

# 최소 크기(_MIN_IMAGE_BYTES) 를 넘기는 PNG — 뒤에 패딩을 붙여 부피만 키운다.
_PNG_HEADER = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4"
    "890000000a49444154789c6300010000050001"
    "0d0a2db40000000049454e44ae426082")


def _png(size: int = 20 * 1024) -> bytes:
    return _PNG_HEADER + b"\x00" * max(0, size - len(_PNG_HEADER))


def _make_docx(path, *, with_image: bool, image_bytes: bytes | None = None):
    d = docx.Document()
    d.add_heading("연차 규정", level=1)
    d.add_paragraph("연차는 15일입니다.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "구분"
    t.cell(0, 1).text = "일수"
    t.cell(1, 0).text = "연차"
    t.cell(1, 1).text = "15"
    if with_image:
        img = path.parent / "shot.png"
        img.write_bytes(image_bytes if image_bytes is not None else _png())
        d.add_picture(str(img))
    d.save(path)


def test_text_and_table_are_extracted(tmp_path):
    p = tmp_path / "규정.docx"
    _make_docx(p, with_image=False)
    res = DocxParser().parse(str(p))
    types = {e.element_type for e in res.elements}
    assert ChunkType.TEXT in types and ChunkType.TABLE in types


def test_embedded_image_is_ocred(tmp_path):
    p = tmp_path / "캡처포함.docx"
    _make_docx(p, with_image=True)
    res = DocxParser(image_ocr=lambda blob: "이미지 속 표: 상여금 지급률 200%").parse(str(p))
    captions = [e for e in res.elements if e.element_type == ChunkType.IMAGE_CAPTION]
    assert len(captions) == 1
    assert "상여금 지급률 200%" in captions[0].text


def test_image_ignored_without_ocr_callback(tmp_path):
    """OCR 미설정이면 이미지는 건너뛴다(텍스트 경로는 그대로 동작)."""
    p = tmp_path / "캡처포함.docx"
    _make_docx(p, with_image=True)
    res = DocxParser().parse(str(p))
    assert not [e for e in res.elements if e.element_type == ChunkType.IMAGE_CAPTION]


def test_tiny_images_are_skipped(tmp_path):
    """로고·아이콘 같은 작은 장식 이미지는 OCR 하지 않는다."""
    p = tmp_path / "로고.docx"
    _make_docx(p, with_image=True, image_bytes=_png(100))
    calls = []

    def ocr(blob):
        calls.append(blob)
        return "로고"

    res = DocxParser(image_ocr=ocr).parse(str(p))
    assert calls == []
    assert not [e for e in res.elements if e.element_type == ChunkType.IMAGE_CAPTION]


def test_ocr_failure_does_not_break_parsing(tmp_path):
    p = tmp_path / "캡처포함.docx"
    _make_docx(p, with_image=True)

    def boom(blob):
        raise RuntimeError("OCR 서비스 다운")

    res = DocxParser(image_ocr=boom).parse(str(p))
    assert [e for e in res.elements if e.element_type == ChunkType.TEXT]


def test_registry_wires_ocr_into_docx_parser():
    """레지스트리에서 docx 파서에도 OCR 콜백이 주입되어야 한다."""
    parser = get_parser(FileFormat.DOCX, ocr=lambda blob: "x")
    assert parser._image_ocr is not None
    assert get_parser(FileFormat.DOCX)._image_ocr is None
