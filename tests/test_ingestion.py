"""적재 파이프라인 통합 테스트 (외부 서비스 없이 fake 주입)."""

from pathlib import Path

import docx
import openpyxl
import pytest

from app.ingestion.chunking import chunk_elements
from app.ingestion.enrichment import enrich
from app.ingestion.intake import DuplicateError, detect_format, intake
from app.ingestion.parsers import get_parser
from app.ingestion.parsers.base import ParsedElement
from app.ingestion.pipeline import apply_review, index, run_auto_stages
from app.schemas.enums import ChunkType, DocStatus, DocType, FileFormat
from app.schemas.ingestion import IngestionStatus
from app.schemas.metadata import DocumentMetadata, GovernanceBlock, IdentificationBlock


# ── 파서 ─────────────────────────────────────────────────────────────────────
def test_txt_parser(tmp_path: Path):
    p = tmp_path / "a.txt"
    p.write_text("첫 문단.\n\n둘째 문단.", encoding="utf-8")
    result = get_parser(FileFormat.TXT).parse(str(p))
    assert [e.text for e in result.elements] == ["첫 문단.", "둘째 문단."]


def test_docx_parser_sections_and_tables(tmp_path: Path):
    d = docx.Document()
    d.add_heading("인사 규정", level=1)
    d.add_paragraph("연차는 15일이다.")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "구분"
    t.rows[0].cells[1].text = "일수"
    t.rows[1].cells[0].text = "연차"
    t.rows[1].cells[1].text = "15"
    p = tmp_path / "policy.docx"
    d.save(str(p))

    result = get_parser(FileFormat.DOCX).parse(str(p))
    types = [e.element_type for e in result.elements]
    assert ChunkType.TABLE in types
    heading_el = next(e for e in result.elements if e.text == "인사 규정")
    assert heading_el.section_title == "인사 규정"
    table_el = next(e for e in result.elements if e.element_type == ChunkType.TABLE)
    assert "구분" in table_el.text and "| --- |" in table_el.text


def test_xlsx_parser_serializes_sheet_with_header(tmp_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "급여"
    ws.append(["사번", "월급"])
    ws.append(["1001", "5000000"])
    p = tmp_path / "salary.xlsx"
    wb.save(str(p))

    result = get_parser(FileFormat.XLSX).parse(str(p))
    assert len(result.elements) == 1
    el = result.elements[0]
    assert el.element_type == ChunkType.TABLE
    assert el.section_title == "급여"
    assert "사번" in el.text and "1001" in el.text


# ── 청킹 ─────────────────────────────────────────────────────────────────────
def test_chunking_keeps_tables_whole_and_splits_long_text():
    long_text = "문장. " * 500  # 매우 긴 텍스트
    els = [
        ParsedElement(text=long_text, element_type=ChunkType.TEXT),
        ParsedElement(text="| a | b |\n| --- | --- |\n| 1 | 2 |",
                      element_type=ChunkType.TABLE),
    ]
    chunks = chunk_elements("doc-x", els)
    text_chunks = [c for c in chunks if c.meta.chunk_type == ChunkType.TEXT]
    table_chunks = [c for c in chunks if c.meta.chunk_type == ChunkType.TABLE]
    assert len(text_chunks) > 1           # 긴 텍스트는 분할됨
    assert len(table_chunks) == 1         # 표는 통째로
    assert all(c.meta.chunk_id.startswith("doc-x::") for c in chunks)


# ── intake / 중복 ───────────────────────────────────────────────────────────
def test_detect_format():
    assert detect_format("a.PDF") == FileFormat.PDF
    assert detect_format("b.jpeg") == FileFormat.JPG
    assert detect_format("c.eml") == FileFormat.EMAIL
    with pytest.raises(ValueError):
        detect_format("c.zip")


def test_email_parser(tmp_path: Path):
    p = tmp_path / "m.eml"
    p.write_text("From: a@co.com\nTo: b@co.com\nSubject: 연차 안내\n\n"
                 "연차는 15일입니다. 확인 바랍니다.", encoding="utf-8")
    result = get_parser(FileFormat.EMAIL).parse(str(p))
    texts = " ".join(e.text for e in result.elements)
    assert "연차 안내" in texts and "15일" in texts


def test_intake_detects_duplicate(tmp_path: Path):
    p = tmp_path / "dup.txt"
    p.write_text("내용", encoding="utf-8")
    with pytest.raises(DuplicateError):
        intake(str(p), ingested_by="admin", hash_lookup=lambda h: "existing-doc")


# ── enrichment (fake LLM) ───────────────────────────────────────────────────
class FakeLLM:
    def __init__(self, response):
        self.response = response

    def complete_json(self, prompt, schema):
        return self.response


def _ident(doc_id="doc-1", fmt=FileFormat.TXT):
    from datetime import datetime, timezone
    return IdentificationBlock(
        doc_id=doc_id, source_filename="f.txt", file_format=fmt,
        file_hash="h", ingested_at=datetime.now(timezone.utc), ingested_by="admin")


def test_enrich_fills_high_conf_and_skips_low_conf():
    doc = DocumentMetadata(identification=_ident())
    llm = FakeLLM({
        "doc_type": {"value": "report", "confidence": 0.95},
        "language": {"value": "ko", "confidence": 0.99},
        "title_normalized": {"value": "연차 규정", "confidence": 0.9},
        "department": {"value": "인사팀", "confidence": 0.3},  # 낮음 → 미채움
        "summary": {"value": "연차 규정 요약", "confidence": 0.9},
        "keywords": ["연차", "휴가"],
        "expected_qa": [{"question": "연차 며칠?", "answer": "15일"}],
        "status": {"value": "active", "confidence": 0.8},
    })
    enrich(doc, "연차는 15일", client=llm, model_name="qwen3.6-27b")
    assert doc.classification.doc_type == DocType.REPORT
    assert doc.classification.title_normalized == "연차 규정"
    assert doc.classification.department is None            # 낮은 confidence
    assert doc.classification.keywords == ["연차", "휴가"]
    assert doc.classification.expected_qa == [{"question": "연차 며칠?", "answer": "15일"}]
    assert doc.lifecycle.status == DocStatus.ACTIVE
    # auto_filled 기록에는 낮은 confidence 필드도 남는다
    fields = {a.field for a in doc.provenance.auto_filled}
    assert "department" in fields and "doc_type" in fields


# ── 전체 파이프라인 (fake 주입) ─────────────────────────────────────────────
class FakeEmbedder:
    def embed(self, texts):
        return [[0.0] * 4 for _ in texts]


class FakeIndexer:
    def __init__(self):
        self.upserted = None

    def upsert(self, vectors, payloads, ids):
        self.upserted = (vectors, payloads, ids)


_GOOD_LLM = FakeLLM({
    "doc_type": {"value": "report", "confidence": 0.95},
    "language": {"value": "ko", "confidence": 0.99},
    "summary": {"value": "요약", "confidence": 0.9},
    "keywords": ["연차"],
    "expected_qa": [{"question": "연차?", "answer": "15일"}],
    "status": {"value": "active", "confidence": 0.9},
})


def test_full_pipeline_happy_path(tmp_path: Path):
    p = tmp_path / "policy.txt"
    p.write_text("연차는 15일입니다.\n\n병가는 별도.", encoding="utf-8")

    ctx = run_auto_stages(str(p), ingested_by="admin",
                          llm=_GOOD_LLM, llm_model="qwen3.6-27b")
    assert ctx.status == IngestionStatus.PENDING_REVIEW
    assert len(ctx.chunks) == 2

    # 접근 토큰은 조직도에서 확장되지만, 직접 apply_review 테스트는 토큰을 직접 지정.
    gov = GovernanceBlock(access_tokens=["n:1"], author_name="hr.manager")
    result = apply_review(ctx, governance=gov,
                          lifecycle_overrides={"status": DocStatus.ACTIVE})
    assert result.ok
    assert ctx.status == IngestionStatus.VALIDATED

    embedder, indexer = FakeEmbedder(), FakeIndexer()
    n = index(ctx, embedder=embedder, indexer=indexer)
    assert n == 2
    assert ctx.status == IngestionStatus.INDEXED
    # payload가 접근통제 토큰을 상속했는지
    _, payloads, _ = indexer.upserted
    assert payloads[0]["access_groups"] == ["n:1"]


def test_read_error_when_ai_cannot_fill_mandatory(tmp_path: Path):
    from app.ingestion.enrichment import ReadError
    p = tmp_path / "scan.txt"
    p.write_text("x", encoding="utf-8")
    # summary/keywords/expected_qa 를 못 주는 LLM → 파일 읽기 실패로 판단
    poor = FakeLLM({"doc_type": {"value": "unknown", "confidence": 0.2},
                    "language": {"value": "unknown", "confidence": 0.2}})
    with pytest.raises(ReadError):
        run_auto_stages(str(p), ingested_by="a", llm=poor, llm_model="m")


def test_full_pipeline_blocks_on_draft_status(tmp_path: Path):
    p = tmp_path / "doc.txt"
    p.write_text("내용입니다.", encoding="utf-8")
    ctx = run_auto_stages(str(p), ingested_by="admin",
                          llm=_GOOD_LLM, llm_model="qwen3.6-27b")

    # draft 상태로는 색인 불가 → 차단
    gov = GovernanceBlock(access_tokens=["n:1"])
    result = apply_review(ctx, governance=gov,
                          lifecycle_overrides={"status": DocStatus.DRAFT})
    assert not result.ok
    assert ctx.status == IngestionStatus.BLOCKED

    # 색인 시도 시 차단 상태라 거부
    with pytest.raises(ValueError):
        index(ctx, embedder=FakeEmbedder(), indexer=FakeIndexer())


def test_blocked_then_corrected_indexes(tmp_path: Path):
    p = tmp_path / "doc.txt"
    p.write_text("내용입니다.", encoding="utf-8")
    ctx = run_auto_stages(str(p), ingested_by="admin",
                          llm=_GOOD_LLM, llm_model="qwen3.6-27b")

    apply_review(ctx, governance=GovernanceBlock(),
                 lifecycle_overrides={"status": DocStatus.DRAFT})
    assert ctx.status == IngestionStatus.BLOCKED

    result = apply_review(ctx, governance=GovernanceBlock(access_tokens=["n:1"]),
                          lifecycle_overrides={"status": DocStatus.ACTIVE})
    assert result.ok
    assert ctx.status == IngestionStatus.VALIDATED
